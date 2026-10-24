"""Экспорт перспективной сетки в Excel (один лист «Сетка»).

Формат ячейки: первая строка — `kinopoisk_id` цифрой, вторая — название
(для серии добавляется «Сезон N, Серия M»). Нет в базе — красный шрифт,
повтор — светло-серая заливка.

Данные берутся на сервере из БД Planner/Oplan3 (`repositories.fetch_window`),
браузер только запрашивает файл.
"""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schedule_perspective import constants

_WEEKDAYS = ('Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс')

_MISSING_FONT = Font(color='CC0000')                    # красный — нет в базе
_REPEAT_FILL = PatternFill('solid', fgColor='EDEDED')   # светло-серый — повтор
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill('solid', fgColor='F2F2F2')
_TITLE_FONT = Font(bold=True, size=13)
_CELL_ALIGN = Alignment(wrap_text=True, vertical='top', horizontal='left')
_CENTER_ALIGN = Alignment(horizontal='center', vertical='top')

_ROW_HEIGHT = 44
_DAY_COLUMN_WIDTH = 30


def _slot_key(slot):
    return (
        slot.get('oplan_program_id') or slot.get('kinopoisk_id'),
        slot.get('season_num'),
        slot.get('episode_num'),
    )


def _repeat_counts(slots):
    counts = {}
    for slot in slots:
        key = _slot_key(slot)
        counts[key] = counts.get(key, 0) + 1
    return counts


def _label(slot):
    """Читаемое название; для серии — с номером сезона/серии."""
    title = (slot.get('title') or '').strip() or 'Без названия'
    episode = slot.get('episode_num')
    if episode is None:
        return title
    season = slot.get('season_num')
    if season is None:
        return f'{title} (Серия {episode})'
    return f'{title} (Сезон {season}, Серия {episode})'


def _cell_value(slot):
    kp_id = slot.get('kinopoisk_id')
    label = _label(slot)
    return f'{kp_id}\n{label}' if kp_id else label


def build_schedule_workbook(channel, start, days, slots):
    """Собирает .xlsx (bytes) из слотов окна."""
    end = start + timedelta(days=days - 1)
    day_dates = [start + timedelta(days=i) for i in range(days)]
    counts = _repeat_counts(slots)

    by_day = {d: {} for d in day_dates}
    positions = set()
    for slot in slots:
        try:
            day = date.fromisoformat(slot['day_date'])
        except (KeyError, TypeError, ValueError):
            continue
        if day not in by_day:
            continue
        by_day[day][slot['position']] = slot
        positions.add(slot['position'])
    position_list = sorted(positions)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Сетка'

    ncols = 1 + len(day_dates)
    channel_name = constants.CHANNELS.get(channel, str(channel))
    title = f'Перспективная сетка — {channel_name} — {start:%d.%m.%Y}–{end:%d.%m.%Y}'
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)

    ws.cell(row=2, column=1, value='№').font = _HEADER_FONT
    ws.cell(row=2, column=1).fill = _HEADER_FILL
    for idx, day in enumerate(day_dates):
        header = ws.cell(row=2, column=2 + idx,
                         value=f'{_WEEKDAYS[day.weekday()]} {day:%d.%m.%Y}')
        header.font = _HEADER_FONT
        header.fill = _HEADER_FILL
        header.alignment = Alignment(horizontal='center', vertical='center')

    for row, position in enumerate(position_list, start=3):
        pos_cell = ws.cell(row=row, column=1, value=position)
        pos_cell.font = _HEADER_FONT
        pos_cell.alignment = _CENTER_ALIGN
        ws.row_dimensions[row].height = _ROW_HEIGHT
        for idx, day in enumerate(day_dates):
            slot = by_day[day].get(position)
            if not slot:
                continue
            cell = ws.cell(row=row, column=2 + idx, value=_cell_value(slot))
            cell.alignment = _CELL_ALIGN
            if counts.get(_slot_key(slot), 0) > 1:
                cell.fill = _REPEAT_FILL
            if not slot.get('in_base'):
                cell.font = _MISSING_FONT

    ws.column_dimensions['A'].width = 5
    for idx in range(len(day_dates)):
        ws.column_dimensions[get_column_letter(2 + idx)].width = _DAY_COLUMN_WIDTH
    ws.freeze_panes = 'B3'

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
